"""Deterministic rule engine: compares draft content with approved facts."""
import io, re
from datetime import date

def label(mo): return f"{mo//12}-year" if mo % 12 == 0 else f"{mo}-month"

# ---- phrasing helpers: number words, fractions and "for life" wording --------------------------
_ONES = {w: i for i, w in enumerate("zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen seventeen eighteen nineteen".split())}
_TENS = {w: 10 * i for i, w in enumerate("_ _ twenty thirty forty fifty sixty seventy eighty ninety".split()) if w != "_"}
_ONE_W = "|".join(sorted(_ONES, key=len, reverse=True))
NUMW = r"(?:(?:" + "|".join(_TENS) + r")(?:[\s-](?:one|two|three|four|five|six|seven|eight|nine))?|" + _ONE_W + r"|(?:one\s+|a\s+)?hundred)"

def num(s):
    """'15', 'fifteen', 'twenty-five', 'a hundred' -> int."""
    s = s.strip().lower()
    if s.isdigit(): return int(s)
    if s.endswith("hundred"): return 100
    parts = re.split(r"[\s-]+", s)
    return sum(_TENS.get(p, _ONES.get(p, 0)) for p in parts)

FRACTIONS = {"half": 50, "1/2": 50, "third": 33.3, "quarter": 25}
LIFE_NOUN = r"(?:warranty|guarantee|coverage|protection|replacements?)"
LIFE = (r"(?:(?:lifetime|life[\s-]?long|forever|never[\s-]expiring)[\s-]+" + LIFE_NOUN + r"\b"                      # lifetime warranty
        r"|" + LIFE_NOUN + r"\s+(?:for\s+(?:life|a\s+lifetime|the\s+lifetime|your\s+lifetime|ever)\b"          # warranty for life
        r"|(?:that|which)\s+(?:lasts?|stays?)\s+(?:forever|for\s+life|a\s+lifetime)"                           # warranty that lasts forever
        r"|(?:never\s+expires?|for\s+the\s+life\s+of\s+(?:the|your)\s+\w+)|forever\b|valid\s+for\s+life))")    # warranty never expires
DUR = r"(\d+(?:\.\d+)?|" + NUMW + r")[\s-]*(year|yr|month)s?"
WARRANTY_RE = re.compile(LIFE + r"|" + DUR + r"[\s-]*(?:limited\s+)?(?:warranty|guarantee)\b"
                         r"|(?:warranty|guarantee)\b\s*(?:period\s*)?(?:of|for|:|-|\u2013|\u2014)?\s*" + DUR, re.I)
DISCOUNT_RE = re.compile(
    r"(?P<pre>\b(?:save|flat|get|up to)\s+)?"
    r"(?:(?P<n>\d{1,3}|" + NUMW + r")\s*(?:%|per[\s-]?cent)"                                                   # 50% / fifty percent
    r"|(?:(?P<fr>half|1/2|a\s+third|one\s+third|a\s+quarter|one\s+quarter)(?:[\s-]*(?:price[d]?|off)\b|\s+the\s+price\b)))"  # half price / a third off
    r"(?:\s*(?:off|discount)\b)?", re.I)
BOGO_RE = re.compile(r"\bbogo\b|\bbuy\s+(?:one|1)[\s,]+get\s+(?:one|1)(?:\s+free)?|\b(?:2|two)\s+for\s+(?:1|one)\b|\b(?:two|2)\s+for\s+the\s+price\s+of\s+(?:one|1)\b", re.I)

def _fraction_value(fr):
    f = fr.lower()
    return 50 if f in ("half", "1/2") else 33.3 if "third" in f else 25

VERIFIED = "Verified against approved evidence"
MISMATCH = "Mismatch with approved information"
UNSUPPORTED = "Unsupported claim"
HUMAN_REVIEW = "Needs human review"


def _fact_evidence(fact):
    return {"fact_id": fact.get("id"), "kind": fact.get("kind", ""), "subject": fact.get("subject", ""),
            "value": fact.get("value", ""), "source_document_id": fact.get("source_document_id"),
            "source_document": fact.get("source", ""), "source_quote": fact.get("source_quote", ""),
            "fact_version": fact.get("version", 1), "valid_from": fact.get("valid_from", ""),
            "valid_to": fact.get("valid_to", "")}


def _number(value):
    try:
        value = float(str(value).replace(",", ""))
        return value if value >= 0 and value < float("inf") else None
    except (TypeError, ValueError):
        return None


def check(text, facts, today=None, include_verified=False):
    today = today or date.today().isoformat()
    ok = lambda f: (not f.get("valid_from") or f["valid_from"] <= today) and (not f.get("valid_to") or today <= f["valid_to"])
    allf = lambda k: [f for f in facts if f["kind"] == k]
    live = lambda k: [f for f in allf(k) if ok(f)]
    out = []
    def add(typ, cat, span, det, appr, src, why, fix=None, status="issue", classification=None, evidence=()):
        finding = dict(type=typ, category=cat, phrase=text[span[0]:span[1]], start=span[0], end=span[1],
                       detected=det, approved=appr, source=src, reason=why, suggestion=fix, status=status)
        finding["classification"] = classification or (MISMATCH if status == "issue" else HUMAN_REVIEW)
        if evidence:
            snapshots = [_fact_evidence(fact) for fact in evidence]
            finding["evidence"] = snapshots
            if len(snapshots) == 1:
                finding.update(snapshots[0])
                finding["source"] = snapshots[0]["source_document"] or src
        out.append(finding)

    def fact_sources(records):
        sources = list(dict.fromkeys(f.get("source", "") for f in records if f.get("source")))
        return sources[0] if len(sources) == 1 else "-"

    def conflict(records, key):
        values = [key(f) for f in records]
        return len(set(values)) > 1

    def decision_fact_finding(typ, cat, span, detected, value, records, value_key, mismatch_type,
                              fix=None, mismatch_reason=None, mismatch_status="issue",
                              mismatch_classification=MISMATCH, tolerance=0.5):
        parsed = [(fact, value_key(fact)) for fact in records]
        if any(value is None for _, value in parsed) or len({value for _, value in parsed}) > 1:
            shown_values = list(dict.fromkeys(str(value) for _, value in parsed if value is not None))
            approved = "Conflicting approved records" + (f": {', '.join(shown_values)}" if shown_values else "")
            add("Conflicting approved facts", cat, span, detected, approved, fact_sources(records),
                "Applicable approved facts conflict or contain an invalid value; a reviewer must resolve them.",
                status="needs_review", classification=HUMAN_REVIEW, evidence=records)
            return
        fact, approved_value = parsed[0]
        if abs(approved_value - value) >= tolerance:
            add(mismatch_type, cat, span, detected, str(fact["value"]), fact_sources(records),
                mismatch_reason(fact) if mismatch_reason else "The claim does not match the approved fact.",
                fix(fact) if fix else None, status=mismatch_status,
                classification=mismatch_classification, evidence=records)
        elif include_verified:
            add("Claim verified", cat, span, detected, str(fact["value"]), fact_sources(records),
                "The claim matches the applicable approved fact.", status="verified",
                classification=VERIFIED, evidence=records)

    # discounts: "50% off", "fifty percent off", "half price", "a third off"
    for m in DISCOUNT_RE.finditer(text):
        if not re.search(r"save|flat|get|up to|off|discount|price", m.group(0), re.I): continue
        word = m.group("fr") is not None
        v = _fraction_value(m.group("fr")) if word else num(m.group("n"))
        if not word and not 0 < v <= 100: continue
        shown = (f"{v:g}" if word else str(v)) + "%"
        L, expired = live("discount"), [f for f in allf("discount") if not ok(f)]
        fix = lambda a: f"{a['value']}% off" if word or not m.group("n").isdigit() else m.group(0).replace(m.group("n"), a["value"])
        if L:
            decision_fact_finding("Discount mismatch", "Offer mismatches", m.span(), shown, v, L,
                                  lambda f: _number(f.get("value")), "Discount mismatch", fix,
                                  lambda f: f"The approved offer is {f['value']}%.")
        elif expired:
            values = [_number(f.get("value")) for f in expired]
            if any(value is None for value in values) or len(set(values)) > 1:
                add("Conflicting approved facts", "Offer mismatches", m.span(), shown,
                    "Expired or conflicting offer records", fact_sources(expired),
                    "No current offer can be selected from the available records; a reviewer must resolve validity or conflicts.",
                    status="needs_review", classification=HUMAN_REVIEW, evidence=expired)
            else:
                a = expired[0]
                add("Offer not valid today", "Offer mismatches", m.span(), shown,
                    f"{a['value']}% (valid through {a.get('valid_to') or 'not recorded'})", fact_sources(expired),
                    "The approved promotion is outside its validity dates.",
                    status="issue", classification=MISMATCH, evidence=expired)
        elif not L:
            add("No approved offer", "Unsupported claims", m.span(), shown, "none on record", "-",
                "No approved offer supports this discount; this absence does not prove the claim is false.",
                status="needs_review", classification=UNSUPPORTED)
    for m in BOGO_RE.finditer(text):  # "buy one get one free" is a different kind of offer: needs an approved record
        add("No approved offer", "Unsupported claims", m.span(), m.group(0), "none on record", "-",
            "No approved offer supports this multi-buy deal; this absence does not prove the claim is false.",
            status="needs_review", classification=UNSUPPORTED)
    # warranty: "1-year warranty", "twelve month guarantee", "lifetime warranty", "warranty for life", "protection for life"
    for m in WARRANTY_RE.finditer(text):
        g = m.groups()
        if g[0] is None and g[2] is None:  # a "for life" phrasing
            mo, noun = 10**6, re.search(LIFE_NOUN, m.group(0), re.I).group(0).lower()
        else:
            n, u = (g[0], g[1]) if g[0] is not None else (g[2], g[3])
            mo = round(float(n) if re.match(r"[\d.]+$", n) else num(n)) * (1 if u.lower().startswith("month") else 12); noun = "warranty"
        hard = noun in ("warranty", "guarantee")  # looser nouns (protection, coverage, replacements) are never auto-rewritten
        W, expired, det = live("warranty"), [f for f in allf("warranty") if not ok(f)], ("lifetime" if mo >= 10**6 else label(mo))
        if not W:
            if expired:
                add("Warranty validity unclear", "Unsupported claims", m.span(), det, "No currently valid record",
                    fact_sources(expired), "The available warranty record is outside its validity dates; a reviewer must confirm the current terms.",
                    status="needs_review", classification=HUMAN_REVIEW, evidence=expired)
            else:
                add("Warranty unsupported", "Unsupported claims", m.span(), det, "none on record", "-",
                    "No approved warranty record exists; this absence does not prove the claim is false.",
                    status="needs_review", classification=UNSUPPORTED)
        else:
            decision_fact_finding("Warranty unsupported", "Unsupported claims", m.span(), det, mo, W,
                                      lambda f: _number(f.get("value")), "Warranty unsupported",
                                  (lambda f: f"{label(int(f['value']))} warranty") if hard else None,
                                      lambda f: f"The approved warranty period is {f['value']} months."
                                      if hard else "This wording may imply broader protection than the approved warranty; a reviewer must interpret it.",
                                      mismatch_status="issue" if hard else "needs_review",
                                      mismatch_classification=MISMATCH if hard else HUMAN_REVIEW, tolerance=0.01)
    # prices
    for m in re.finditer(r"(?:₹|rs\.?\s*|\$\s*)\s*(?P<amount>[\d,]+(?:\.\d{1,2})?)", text, re.I):
        sentence_start = max(text.rfind(mark, 0, m.start()) for mark in ".!?\n;") + 1
        ends = [text.find(mark, m.end()) for mark in ".!?\n;"]
        sentence_end = min((end for end in ends if end >= 0), default=len(text))
        sentence = text[sentence_start:sentence_end]
        amount = _number(m.group("amount"))
        if amount is None:
            continue
        approved_subjects = [f["subject"] for f in allf("price") if f.get("subject")]
        matched_subjects = [subject for subject in approved_subjects if re.search(r"\b" + re.escape(subject) + r"\b", sentence, re.I)]
        subject = max(matched_subjects, key=len) if matched_subjects else ""
        applicable = [f for f in live("price") if subject and f["subject"].casefold() == subject.casefold()]
        expired = [f for f in allf("price") if subject and f["subject"].casefold() == subject.casefold() and not ok(f)]
        has_price_context = bool(re.search(r"\b(?:price|cost|priced|only|just)\b", sentence, re.I))
        span = m.span("amount")
        if applicable:
            decision_fact_finding("Price mismatch", "Offer mismatches", span, m.group("amount"), amount, applicable,
                                  lambda f: _number(f.get("value")), "Price mismatch", lambda f: str(f["value"]),
                                  lambda f: f"The approved price of {subject} is {f['value']}.", tolerance=0.005)
        elif expired:
            add("Price validity unclear", "Offer mismatches", span, m.group("amount"), "No currently valid price",
                fact_sources(expired), "The approved price record is outside its validity dates; a reviewer must confirm the current price.",
                status="needs_review", classification=HUMAN_REVIEW, evidence=expired)
        elif has_price_context:
            add("Price unsupported", "Unsupported claims", span, m.group("amount"), "no approved price record", "-",
                "No approved price record matches this claim; this absence does not prove the price is false.",
                status="needs_review", classification=UNSUPPORTED)
    # wording
    for k, typ in (("banned", "Banned phrase"), ("term", "Naming drift"), ("tone", "Tone keyword")):
        groups = {}
        for fact in live(k):
            groups.setdefault(fact["subject"].casefold(), []).append(fact)
        for records in groups.values():
            fact = records[0]
            for m in re.finditer(r"\b" + re.escape(fact["subject"]) + r"\b", text, re.I):
                values = {record.get("value", "") for record in records}
                if len(values) > 1:
                    add("Conflicting approved facts", "Tone / naming", m.span(), m.group(0),
                        "Conflicting approved terminology", fact_sources(records),
                        "Applicable approved terminology records conflict; a reviewer must resolve them.",
                        status="needs_review", classification=HUMAN_REVIEW, evidence=records)
                    continue
                why = f"'{fact['subject']}' is not allowed." if k == "banned" else f"Use '{fact['value']}' instead."
                add(typ, "Tone / naming", m.span(), m.group(0), fact["value"] or "remove", fact_sources(records), why,
                    fact["value"] or None, evidence=records, classification=MISMATCH)
    for m in re.finditer(r"!{2,}", text):
        add("Tone keyword", "Tone / naming", m.span(), m.group(0), "!", "Tone guide", "Use at most one exclamation mark.", "!")
    # Conservative brand-voice heuristics: human-review prompts, not claims of factual falsity.
    for pattern, title, explanation, suggestion in (
        (r"\b(?:act now|buy now|hurry|last chance|don't miss out|do not miss out)\b", "High-pressure sales language",
         "This phrase can create urgency or pressure. Check whether the approved brand voice permits this style.",
         "Review the urgency wording against the approved brand voice."),
        (r"\b(?:unbeatable|unmatched|ultimate|perfect|revolutionary)\b", "Promotional superlative",
         "This is a subjective promotional superlative. Confirm that the wording is approved and appropriately substantiated.",
         None),
    ):
        for m in re.finditer(pattern, text, re.I):
            sentence_start = max(text.rfind(".", 0, m.start()), text.rfind("!", 0, m.start()), text.rfind("?", 0, m.start())) + 1
            ends = [position for position in (text.find(".", m.end()), text.find("!", m.end()), text.find("?", m.end())) if position >= 0]
            sentence_end = min(ends) + 1 if ends else len(text)
            sentence_text = text[sentence_start:sentence_end]
            sentence_has_finding = any(item.get("end", 0) > sentence_start and item.get("start", 0) < sentence_end for item in out)
            # Avoid duplicate tone flags in offer/price sentences; claim rules and the reviewer handle those.
            offer_context = bool(re.search(r"\b(?:price|prices|cost|deal|offer|discount|save|off)\b", sentence_text, re.I))
            if sentence_has_finding or offer_context:
                continue
            add(title, "Tone / naming", m.span(), m.group(0), "Use approved, evidence-based language", "Brand voice review",
                explanation, suggestion, status="needs_review", classification=HUMAN_REVIEW)
    # Excessive uppercase can conflict with a calm/professional voice; review rather than auto-reject.
    for m in re.finditer(r"(?<![A-Za-z])[A-Z][A-Z0-9 ,.!?'-]{11,}(?![A-Za-z])", text):
        phrase = m.group(0).strip()
        letters = [c for c in phrase if c.isalpha()]
        overlaps_known_finding = any(item.get("end", 0) > m.start() and item.get("start", 0) < m.end() for item in out)
        if not overlaps_known_finding and len(letters) >= 10 and sum(c.isupper() for c in letters) / len(letters) > .9:
            add("Excessive uppercase", "Tone / naming", m.span(), phrase, "Sentence case unless the brand guide says otherwise",
                "Brand voice review", "This text is predominantly uppercase. Confirm whether the brand style guide permits all-caps copy.",
                phrase[:1].upper() + phrase[1:].lower(), status="needs_review", classification=HUMAN_REVIEW)
    for m in re.finditer(r"click here", text, re.I):
        add("Accessibility cue", "Asset rule issues", m.span(), m.group(0), "descriptive link text", "Accessibility rule",
            "Link text should describe the destination.", "read the details")
    # risky claims need an approved claim record
    for m in re.finditer(r"#1|number one|best[- ]in[- ]class|clinically proven|guaranteed|100%\s*(?:safe|natural|effective)|world'?s (?:best|first)", text, re.I):
        matches = [f for f in live("claim") if f["subject"].casefold() == m.group(0).casefold()]
        if matches:
            values = {f.get("value", "") for f in matches}
            if len(values) > 1:
                add("Conflicting approved facts", "Unsupported claims", m.span(), m.group(0),
                    "Conflicting approved claim records", fact_sources(matches),
                    "Applicable approved claim records conflict; a reviewer must resolve them.",
                    status="needs_review", classification=HUMAN_REVIEW, evidence=matches)
            elif include_verified:
                add("Claim verified", "Unsupported claims", m.span(), m.group(0), matches[0].get("value") or matches[0]["subject"],
                    fact_sources(matches), "The claim exactly matches an approved claim record.", status="verified",
                    classification=VERIFIED, evidence=matches)
        else:
            expired = [f for f in allf("claim") if f["subject"].casefold() == m.group(0).casefold()]
            if expired:
                add("Claim validity unclear", "Unsupported claims", m.span(), m.group(0), "No currently valid record",
                    fact_sources(expired), "The matching approved claim record is outside its validity dates; a reviewer must confirm it.",
                    status="needs_review", classification=HUMAN_REVIEW, evidence=expired)
            else:
                add("Unsupported claim", "Unsupported claims", m.span(), m.group(0), "no approved record", "-",
                    "No approved record supports this claim; this absence does not prove the claim is false.",
                    status="needs_review", classification=UNSUPPORTED)
    # Accessibility checks for HTML snippets pasted into the checker. These are
    # deliberately conservative: inline CSS can be assessed, but external stylesheets
    # and computed browser styles cannot be inferred from source text alone.
    if re.search(r"<\s*(?:html|body|main|article|section|div|p|a|img|h[1-6])\b", text, re.I):
        for tag in re.finditer(r"<img\b[^>]*>", text, re.I | re.S):
            raw_tag = tag.group(0)
            alt = re.search(r"\balt\s*=\s*(['\"])(.*?)\1", raw_tag, re.I | re.S)
            if alt is None:
                out.append(dict(type="Image missing alternative text", category="Accessibility",
                    phrase=raw_tag, start=tag.start(), end=tag.end(), detected="No alt attribute",
                    approved="Meaningful alternative text or an explicitly decorative image", source="WCAG accessibility check",
                    reason="This image has no alt attribute. Add concise alternative text for informative images; use alt=\"\" only when the image is purely decorative.",
                    suggestion=None, status="needs_review", classification=HUMAN_REVIEW))
            elif not alt.group(2).strip():
                out.append(dict(type="Empty image alternative text", category="Accessibility",
                    phrase=raw_tag, start=tag.start(), end=tag.end(), detected="Empty alt text",
                    approved="Meaningful alternative text or confirmed decorative intent", source="WCAG accessibility check",
                    reason="Empty alt text is correct for decorative images, but informative images need a useful text alternative. Confirm the intent.",
                    suggestion=None, status="needs_review", classification=HUMAN_REVIEW))

        # Evaluate only explicit inline foreground/background pairs. This is a
        # screening result, not a substitute for checking rendered/computed styles.
        def rgb(value):
            value = value.strip().lower()
            if re.fullmatch(r"#[0-9a-f]{3}", value):
                value = "#" + "".join(ch * 2 for ch in value[1:])
            if re.fullmatch(r"#[0-9a-f]{6}", value):
                return tuple(int(value[i:i+2], 16) for i in (1, 3, 5))
            match = re.fullmatch(r"rgb\(\s*(\d{1,3})\s*,\s*(\d{1,3})\s*,\s*(\d{1,3})\s*\)", value)
            if match and all(int(v) <= 255 for v in match.groups()):
                return tuple(map(int, match.groups()))
            return None
        def luminance(color):
            channels = [v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4
                        for v in (x / 255 for x in color)]
            return .2126 * channels[0] + .7152 * channels[1] + .0722 * channels[2]
        for tag in re.finditer(r"<([a-z][a-z0-9]*)\b[^>]*>", text, re.I | re.S):
            style = re.search(r"\bstyle\s*=\s*(['\"])(.*?)\1", tag.group(0), re.I | re.S)
            if not style:
                continue
            declarations = {k.strip().lower(): v.strip() for k, v in
                re.findall(r"([a-z-]+)\s*:\s*([^;]+)", style.group(2), re.I)}
            fg, bg = rgb(declarations.get("color", "")), rgb(declarations.get("background-color", ""))
            if fg and bg:
                a, b = sorted((luminance(fg), luminance(bg)), reverse=True)
                ratio = (a + .05) / (b + .05)
                if ratio < 4.5:
                    out.append(dict(type="Possible low text contrast", category="Accessibility",
                        phrase=tag.group(0), start=tag.start(), end=tag.end(), detected=f"{ratio:.2f}:1 contrast",
                        approved="At least 4.5:1 for normal text (WCAG AA)", source="WCAG contrast screening",
                        reason="The inline foreground and background colours fall below the 4.5:1 normal-text target. Large text has a lower 3:1 threshold; verify rendered styles and text size.",
                        suggestion=None, status="needs_review", classification=HUMAN_REVIEW))

        headings = [(int(m.group(1)), m.start(), m.group(0)) for m in re.finditer(r"<h([1-6])\b[^>]*>", text, re.I)]
        for previous, current_heading in zip(headings, headings[1:]):
            if current_heading[0] > previous[0] + 1:
                out.append(dict(type="Heading level skipped", category="Accessibility",
                    phrase=current_heading[2], start=current_heading[1], end=current_heading[1] + len(current_heading[2]),
                    detected=f"h{current_heading[0]} follows h{previous[0]}", approved="Logical heading hierarchy",
                    source="WCAG document structure check",
                    reason="Heading levels should normally progress in order so assistive-technology users can navigate the document structure.",
                    suggestion=None, status="needs_review", classification=HUMAN_REVIEW))

    # Subjective marketing superlatives are surfaced as review prompts, not factual
    # verdicts. Existing explicit claim rules above remain the source of mismatches.
    for m in re.finditer(r"\b(?:best ever|revolutionary|game[- ]changing)\b", text, re.I):
        if any(i["start"] < m.end() and i["end"] > m.start() for i in out):
            continue
        out.append(dict(type="Promotional superlative", category="Tone / naming", phrase=m.group(0),
            start=m.start(), end=m.end(), detected=m.group(0), approved="Evidence-backed, brand-aligned wording",
            source="Promotional-claims heuristic",
            reason="This emphatic phrase may need substantiation or may conflict with the approved brand voice. It is a review prompt, not proof that the claim is false.",
            suggestion=None, status="needs_review", classification=HUMAN_REVIEW))

    return sorted(out, key=lambda i: i["start"])

def apply_fixes(text, issues):
    last = len(text)
    for i in sorted((i for i in issues if i["status"] == "issue" and i["suggestion"] is not None), key=lambda i: -i["start"]):
        if i["end"] <= last:
            text = text[:i["start"]] + i["suggestion"] + text[i["end"]:]; last = i["start"]
    return text

def _docx_xml_text(data):
    """Backup DOCX reader that needs no extra package: reads the paragraphs straight from word/document.xml."""
    import zipfile, xml.etree.ElementTree as ET
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(io.BytesIO(data)) as z: root = ET.fromstring(z.read("word/document.xml"))
    bit = {W + "t": lambda e: e.text or "", W + "tab": lambda e: "\t", W + "br": lambda e: "\n"}
    return "\n".join("".join(bit[e.tag](e) for e in p.iter() if e.tag in bit) for p in root.iter(W + "p"))

MAX_PDF_PAGES = 100

def _pdf_text(data):
    """Selectable text from a PDF, page by page. Scanned PDFs have none (returns empty)."""
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    if r.is_encrypted and not r.decrypt(""): raise ValueError("password-protected PDF")
    return "\n".join((p.extract_text() or "").strip() for p in r.pages[:MAX_PDF_PAGES]).strip()

def read_text(name, data):
    if name.lower().endswith(".pdf"): return _pdf_text(data)
    if name.lower().endswith(".docx"):
        try:
            from docx import Document
            d = Document(io.BytesIO(data))
            return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
        except Exception:  # python-docx missing, wrong "docx" package installed, or an unusual file: use the backup reader
            return _docx_xml_text(data)
    return data.decode("utf-8", "ignore")

MAX_IMAGE_BYTES = 10 * 1024 * 1024
MAX_IMAGE_PIXELS = 12_000_000
MAX_IMAGE_SIDE = 8_000
MIN_IMAGE_WIDTH, MIN_IMAGE_HEIGHT = 600, 315
COLOR_DISTANCE_TOLERANCE = 90  # Euclidean RGB distance, channels 0-255; documented approximate palette tolerance.


class ImageValidationError(Exception):
    pass


class ImageTooLargeError(ImageValidationError):
    pass


def is_image_signature(data):
    return data.startswith(b"\x89PNG\r\n\x1a\n") or data.startswith(b"\xff\xd8\xff")


def load_image(data):
    """Decode only PNG/JPEG within byte, side and pixel bounds; composite transparency onto white."""
    from PIL import Image
    if not isinstance(data, (bytes, bytearray)) or not data:
        raise ImageValidationError("The image file is empty or invalid.")
    if len(data) > MAX_IMAGE_BYTES:
        raise ImageTooLargeError("The image file is larger than 10 MB.")
    try:
        with Image.open(io.BytesIO(data)) as probe:
            image_format = probe.format
            width, height = probe.size
            if image_format not in ("PNG", "JPEG"):
                raise ImageValidationError("Only PNG and JPEG image data are supported.")
            if (width <= 0 or height <= 0 or width > MAX_IMAGE_SIDE or height > MAX_IMAGE_SIDE
                    or width * height > MAX_IMAGE_PIXELS):
                raise ImageTooLargeError("The decoded image dimensions exceed the safe limit.")
            probe.verify()
        with Image.open(io.BytesIO(data)) as source:
            source.load()
            rgba = source.convert("RGBA")
            background = Image.new("RGBA", rgba.size, "white")
            image = Image.alpha_composite(background, rgba).convert("RGB")
        return image, {"format": image_format, "width": width, "height": height, "pixels": width * height}
    except (ImageValidationError, ImageTooLargeError):
        raise
    except Image.DecompressionBombError as exc:
        raise ImageTooLargeError("The decoded image dimensions exceed the safe limit.") from exc
    except (Image.UnidentifiedImageError, OSError, SyntaxError, ValueError) as exc:
        raise ImageValidationError("The image data is malformed or unsupported.") from exc


def inspect_image(data, facts):
    """Return deterministic image findings and a user-facing record of checks actually performed."""
    import math
    import re
    image, info = load_image(data)
    out, checks = [], [
        {"name": "File size", "status": "passed", "detail": f"{len(data)} bytes; maximum {MAX_IMAGE_BYTES} bytes."},
        {"name": "File format", "status": "passed", "detail": f"Decoded image data is {info['format']}."},
    ]

    def add(typ, why, det, appr, status="issue", evidence=()):
        out.append(dict(type=typ, category="Asset rule issues", phrase="image", start=0, end=0, detected=det,
                        approved=appr, source="Asset rules", reason=why, suggestion=None, status=status,
                        classification=MISMATCH if status == "issue" else HUMAN_REVIEW,
                        evidence=[_fact_evidence(f) for f in evidence]))

    checks.append({"name": "Dimensions", "status": "passed", "detail": f"{info['width']} x {info['height']} px; within the {MAX_IMAGE_SIDE}px side and {MAX_IMAGE_PIXELS:,}-pixel decode limits."})
    if info["width"] < MIN_IMAGE_WIDTH or info["height"] < MIN_IMAGE_HEIGHT:
        add("Image too small", f"Images should be at least {MIN_IMAGE_WIDTH}x{MIN_IMAGE_HEIGHT} px.",
            f"{info['width']}x{info['height']}", f"{MIN_IMAGE_WIDTH}x{MIN_IMAGE_HEIGHT} minimum")
        checks[-1]["status"] = "review"
        checks[-1]["detail"] += f" Recommended minimum is {MIN_IMAGE_WIDTH} x {MIN_IMAGE_HEIGHT} px."

    color_facts = [f for f in facts if f.get("kind") == "color"]
    approved = []
    for fact in color_facts:
        raw = str(fact.get("value", "")).strip()
        if re.fullmatch(r"#[0-9a-fA-F]{6}", raw):
            rgb = tuple(int(raw[index:index + 2], 16) for index in (1, 3, 5))
            approved.append((fact, raw.upper(), rgb))
    if not color_facts:
        checks.append({"name": "Brand colours", "status": "unavailable", "detail": "No approved brand colours are configured."})
    elif not approved:
        checks.append({"name": "Brand colours", "status": "review", "detail": "Approved colour records are not valid #RRGGBB values; palette comparison was not run."})
        add("Brand palette unavailable", "Approved brand colours could not be read, so this image was not compared.",
            "not checked", "valid approved colour values", "needs_review", color_facts)
    else:
        sample = image.resize((64, 64)).quantize(colors=5)
        palette = sample.getpalette() or []
        sampled = sorted(sample.getcolors(64 * 64) or [], reverse=True)
        dominant = [tuple(palette[index * 3:index * 3 + 3]) for _, index in sampled[:1]
                    if len(palette[index * 3:index * 3 + 3]) == 3]
        distances = [(math.sqrt(sum((a - b) ** 2 for a, b in zip(pixel, rgb))), pixel, fact, color)
                     for pixel in dominant for fact, color, rgb in approved]
        nearest = min(distances, default=None, key=lambda item: item[0])
        if nearest and nearest[0] <= COLOR_DISTANCE_TOLERANCE:
            distance, pixel, fact, color = nearest
            checks.append({"name": "Brand colours", "status": "passed", "detail": f"Most frequent colour in the 64x64 sample, #{pixel[0]:02X}{pixel[1]:02X}{pixel[2]:02X}, is within RGB distance {distance:.1f} of approved {color}; tolerance {COLOR_DISTANCE_TOLERANCE}."})
        else:
            expected = ", ".join(color for _, color, _ in approved)
            checks.append({"name": "Brand colours", "status": "review", "detail": f"No sampled dominant colour was within RGB distance {COLOR_DISTANCE_TOLERANCE} of the approved palette ({expected})."})
            add("Colour differs from sampled brand palette", "The sampled dominant colours did not match the approved palette within the configured tolerance.",
                "sampled image colours", expected, "needs_review", [fact for fact, _, _ in approved])

    checks.extend([
        {"name": "Logo presence", "status": "unavailable", "detail": "No approved logo reference asset is configured; logo comparison was not run."},
        {"name": "Text readability", "status": "not_performed", "detail": "No reliable local text-size or contrast assessment is configured."},
    ])
    return out, checks, info


def check_image(data, facts):
    """Compatibility wrapper returning only findings for existing callers."""
    return inspect_image(data, facts)[0]
